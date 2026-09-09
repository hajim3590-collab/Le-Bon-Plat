#!/usr/bin/env python3
"""
Script pour modifier des PowerPoint facilement
Usage: python edit_powerpoint.py
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os

print("=" * 60)
print("📊 ÉDITEUR POWERPOINT - Claude Assistant")
print("=" * 60)

# 1. Lister les PowerPoint disponibles
pptx_files = [f for f in os.listdir('.') if f.endswith('.pptx')]

if not pptx_files:
    print("\n❌ Aucun PowerPoint trouvé dans ce dossier!")
    print("Crée d'abord un PowerPoint ou place-le ici.")
    exit(1)

print(f"\n📁 PowerPoint trouvés ({len(pptx_files)}):")
for idx, file in enumerate(pptx_files, 1):
    print(f"  {idx}. {file}")

# 2. Choisir le fichier
choice = input(f"\nQuel PowerPoint veux-tu modifier? (1-{len(pptx_files)}): ").strip()
try:
    file_idx = int(choice) - 1
    if file_idx < 0 or file_idx >= len(pptx_files):
        raise ValueError
    pptx_file = pptx_files[file_idx]
except:
    print("❌ Choix invalide!")
    exit(1)

# 3. Charger le PowerPoint
try:
    prs = Presentation(pptx_file)
    print(f"\n✅ {pptx_file} chargé ({len(prs.slides)} slides)")
except Exception as e:
    print(f"❌ Erreur: {e}")
    exit(1)

# 4. Afficher les slides
print("\n📋 Slides disponibles:")
for idx, slide in enumerate(prs.slides, 1):
    texts = []
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text.strip():
            texts.append(shape.text.strip()[:50])

    preview = " | ".join(texts[:2]) if texts else "[Slide vide]"
    print(f"  Slide {idx}: {preview}...")

# 5. Menu d'actions
print("\n" + "=" * 60)
print("ACTIONS DISPONIBLES:")
print("=" * 60)
print("1. Ajouter du texte à une slide")
print("2. Modifier le titre d'une slide")
print("3. Voir le contenu complet d'une slide")
print("4. Créer une nouvelle slide")
print("5. Exporter en PDF (si LibreOffice disponible)")
print("6. Sauvegarder et quitter")

action = input("\nQuelle action? (1-6): ").strip()

if action == "1":
    slide_num = int(input("Numéro de la slide: ")) - 1
    if 0 <= slide_num < len(prs.slides):
        slide = prs.slides[slide_num]
        text = input("Texte à ajouter: ")

        # Ajouter une textbox
        left = Inches(1)
        top = Inches(4.5)
        width = Inches(8)
        height = Inches(0.8)

        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(44, 62, 80)

        print("✅ Texte ajouté!")
    else:
        print("❌ Slide invalide!")

elif action == "2":
    slide_num = int(input("Numéro de la slide: ")) - 1
    if 0 <= slide_num < len(prs.slides):
        slide = prs.slides[slide_num]
        new_title = input("Nouveau titre: ")

        # Chercher la shape de titre
        for shape in slide.shapes:
            if shape.has_text_frame and "title" in shape.name.lower():
                shape.text = new_title
                print("✅ Titre modifié!")
                break
        else:
            print("⚠️ Pas de titre trouvé sur cette slide")
    else:
        print("❌ Slide invalide!")

elif action == "3":
    slide_num = int(input("Numéro de la slide: ")) - 1
    if 0 <= slide_num < len(prs.slides):
        slide = prs.slides[slide_num]
        print(f"\n📄 Contenu de la Slide {slide_num + 1}:")
        print("-" * 60)

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                print(f"• {shape.text}")
        print("-" * 60)
    else:
        print("❌ Slide invalide!")

elif action == "4":
    slide_layout = prs.slide_layouts[6]  # Blank layout
    new_slide = prs.slides.add_slide(slide_layout)
    print(f"✅ Nouvelle slide créée (Slide {len(prs.slides)})")

elif action == "5":
    print("⏳ Conversion en PDF...")
    print("Note: Nécessite LibreOffice")
    os.system(f'libreoffice --headless --convert-to pdf "{pptx_file}"')
    pdf_file = pptx_file.replace('.pptx', '.pdf')
    print(f"✅ PDF créé: {pdf_file}")

elif action == "6":
    save_file = pptx_file
    backup = pptx_file.replace('.pptx', '_BACKUP.pptx')

    # Sauvegarder backup
    if os.path.exists(pptx_file):
        os.rename(pptx_file, backup)
        print(f"📦 Backup créé: {backup}")

    # Sauvegarder la version modifiée
    prs.save(pptx_file)
    print(f"✅ {pptx_file} sauvegardé!")
    print("\n✨ Prêt pour la présentation!")
    exit(0)

else:
    print("❌ Action invalide!")

# Sauvegarder à la fin
print("\n💾 Sauvegarder les changements? (y/n): ", end="")
if input().lower() == 'y':
    backup = pptx_file.replace('.pptx', '_BACKUP.pptx')
    if os.path.exists(pptx_file):
        os.rename(pptx_file, backup)
        print(f"📦 Backup créé: {backup}")

    prs.save(pptx_file)
    print(f"✅ {pptx_file} sauvegardé!")
else:
    print("⚠️ Changements non sauvegardés")

print("\n✨ À bientôt!")
